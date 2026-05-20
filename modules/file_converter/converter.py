from __future__ import annotations

from io import BytesIO

import pandas as pd
from docx import Document
from pypdf import PdfReader

from services.file_service import build_timestamped_filename, decode_text, suffix_lower


SUPPORTED_CONVERSION_LABELS = ("PDF 转 文本", "Word 转 文本", "Excel 转 CSV", "TXT 转 Word 文档")


def get_conversion_label_for_file(filename: str) -> str | None:
    suffix = suffix_lower(filename)
    if suffix == ".pdf":
        return "PDF 转 文本"
    if suffix == ".docx":
        return "Word 转 文本"
    if suffix in {".xlsx", ".xlsm"}:
        return "Excel 转 CSV"
    if suffix == ".txt":
        return "TXT 转 Word 文档"
    return None


def _convert_pdf_to_txt(file_bytes: bytes, source_name: str) -> tuple[bytes, str]:
    reader = PdfReader(BytesIO(file_bytes))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # pragma: no cover - defensive path
            raise ValueError("PDF 文件已加密或受密码保护，当前版本暂不支持。") from exc

    pages: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        page_text = (page.extract_text() or "").strip()
        if page_text:
            pages.append(f"--- 第 {page_number} 页 ---\n{page_text}")

    if not pages:
        raise ValueError("未从 PDF 中提取到可复制文本。若是扫描版 PDF，请先做 OCR。")

    text = "\n\n".join(pages)
    return text.encode("utf-8"), build_timestamped_filename(source_name, ".txt", prefix="pdf_to_txt")


def _extract_docx_text(document: Document) -> str:
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)

    for table in document.tables:
        for row in table.rows:
            row_text = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(row_text):
                chunks.append("\t".join(row_text))

    return "\n".join(chunks).strip()


def _convert_docx_to_txt(file_bytes: bytes, source_name: str) -> tuple[bytes, str]:
    document = Document(BytesIO(file_bytes))
    text = _extract_docx_text(document)
    if not text:
        raise ValueError("Word 文档中没有可导出的文本。")
    return text.encode("utf-8"), build_timestamped_filename(source_name, ".txt", prefix="word_to_txt")


def _convert_excel_to_csv(file_bytes: bytes, source_name: str) -> tuple[bytes, str]:
    dataframe = pd.read_excel(BytesIO(file_bytes), sheet_name=0, engine="openpyxl")
    csv_text = dataframe.to_csv(index=False)
    return csv_text.encode("utf-8-sig"), build_timestamped_filename(source_name, ".csv", prefix="excel_to_csv")


def _convert_txt_to_docx(file_bytes: bytes, source_name: str) -> tuple[bytes, str]:
    text = decode_text(file_bytes)
    document = Document()
    lines = text.splitlines() or [""]
    for line in lines:
        document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue(), build_timestamped_filename(source_name, ".docx", prefix="txt_to_word")


def convert_file(file_bytes: bytes, source_name: str, conversion_label: str) -> tuple[bytes, str]:
    suffix = suffix_lower(source_name)

    if conversion_label == "PDF 转 文本":
        if suffix != ".pdf":
            raise ValueError("所选文件不是 PDF。")
        return _convert_pdf_to_txt(file_bytes, source_name)

    if conversion_label == "Word 转 文本":
        if suffix != ".docx":
            raise ValueError("所选文件不是 Word(.docx) 文档。")
        return _convert_docx_to_txt(file_bytes, source_name)

    if conversion_label == "Excel 转 CSV":
        if suffix not in {".xlsx", ".xlsm"}:
            raise ValueError("所选文件不是 Excel(.xlsx/.xlsm) 文件。")
        return _convert_excel_to_csv(file_bytes, source_name)

    if conversion_label == "TXT 转 Word 文档":
        if suffix != ".txt":
            raise ValueError("所选文件不是 TXT 文件。")
        return _convert_txt_to_docx(file_bytes, source_name)

    raise ValueError("不支持的转换类型。")
